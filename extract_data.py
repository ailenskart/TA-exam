#!/usr/bin/env python3
"""Convert the TA exam workbook (.xlsx) into data.js for the study app.

Usage:  python3 extract_data.py <workbook.xlsx>

Re-run this whenever a new phase of questions is added to the workbook —
it regenerates data.js, and the app picks the new questions up automatically.
Every sheet except 'Topic Overview' is treated as a question sheet with the
columns: Topic Name | Topic Basics | Q No. | [Passage] | Question |
Option A-D | Correct Answer | Answer Explanation | Memory Cue | Source.
"""
import json
import sys

import openpyxl

SECTION_OF_PREFIX = {"R": "Reasoning", "E": "English", "G": "GK", "C": "Current Affairs"}


def clean(v):
    return str(v).strip() if v is not None else ""


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    topics = []
    questions = []

    for sheet in wb.sheetnames:
        if sheet == "Topic Overview":
            continue
        ws = wb[sheet]
        header = [clean(c) for c in next(ws.iter_rows(max_row=1, values_only=True))]
        has_passage = "Passage" in header
        section = SECTION_OF_PREFIX.get(sheet[0], "Other")

        topic_id = sheet
        topic_name = sheet
        topic_basics = ""
        count = 0

        for row in ws.iter_rows(min_row=2, values_only=True):
            cells = [clean(c) for c in row]
            offset = 1 if has_passage else 0
            passage = cells[3] if has_passage else ""
            q_text = cells[3 + offset]
            if not q_text:
                continue
            answer = cells[8 + offset][:1].upper()
            if answer not in "ABCD":
                continue
            if cells[0]:
                topic_name = cells[0]
            if cells[1] and not topic_basics:
                topic_basics = cells[1]
            count += 1
            questions.append({
                "id": f"{topic_id}#{count}",
                "t": topic_id,
                "p": passage,
                "q": q_text,
                "o": [cells[4 + offset], cells[5 + offset], cells[6 + offset], cells[7 + offset]],
                "a": "ABCD".index(answer),
                "e": cells[9 + offset],
                "m": cells[10 + offset],
            })

        topics.append({
            "id": topic_id,
            "name": topic_name,
            "section": section,
            "basics": topic_basics,
            "count": count,
        })
        print(f"  {sheet}: {count} questions")

    out = "const TA_DATA = " + json.dumps(
        {"topics": topics, "questions": questions}, ensure_ascii=False
    ) + ";\n"
    with open("data.js", "w", encoding="utf-8") as f:
        f.write(out)
    print(f"\nWrote data.js — {len(questions)} questions in {len(topics)} topics")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "Final_V4.xlsx")
